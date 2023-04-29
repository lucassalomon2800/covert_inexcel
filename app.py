import docx
import openpyxl
import os


workbook = openpyxl.Workbook()
sheet = workbook.active
import os

def realizar():
    folder_path = 'documents'
    for filename in os.listdir(folder_path):
        if filename.endswith('.docx'):
            doc = docx.Document(os.path.join(folder_path, filename))
            data = []
            for paragraph in doc.paragraphs:
                data.append(paragraph.text)
            sheet.append(data)
    workbook.save('example1.xlsx')




from flask import Flask, request, redirect, url_for

app = Flask(__name__)

@app.route("/")
def index():
    realizar()
    with open("index.html", "r") as f:
        return f.read()


@app.route("/documents", methods=["POST"])
def documents():
    file = request.files["file"]
    
    extension = os.path.splitext(file.filename)[1]
    if extension==".doc" or extension==".docx":
        file.save(os.path.join("documents",file.filename))
        message=str("Perfect is upload")
        return redirect(url_for("index"))
    else:
        message=str("error")
        return redirect(url_for("index"))
    
    


if __name__ == "__main__":
    app.run()